"""Excel statement parser (.xlsx via openpyxl, .xls via xlrd).

Also copes with HTML tables that are exported with an .xls extension.
"""

from __future__ import annotations

import html as html_lib
import io
import re
from datetime import date, datetime
from typing import Any

from app.parsers.base import (
    ParsedStatement,
    ParserError,
    sniff_account_details,
    sniff_bank_name,
    sniff_opening_balance,
)
from app.parsers.detector import rows_to_parsed


def _cell_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    return value


def _rows_from_matrix(matrix: list[list[Any]]) -> list[list[Any]]:
    return [[_cell_value(cell) for cell in row] for row in matrix]


def _parse_html_xls(data: bytes) -> list[list[Any]]:
    """Extract table cells from HTML/XML masquerading as .xls."""
    text = data.decode("utf-8", errors="replace")
    matrix: list[list[Any]] = []
    for row_html in re.findall(r"<tr[^>]*>(.*?)</tr>", text, re.I | re.S):
        cells = re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", row_html, re.I | re.S)
        cleaned = []
        for cell in cells:
            plain = re.sub(r"<[^>]+>", "", cell)
            cleaned.append(html_lib.unescape(plain).strip())
        if cleaned:
            matrix.append(cleaned)
    return matrix


def _parse_xlsx(data: bytes) -> list[list[list[Any]]]:
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True
        )
    except Exception as error:  # openpyxl raises many exception types
        raise ParserError("The Excel file could not be opened.") from error

    sheets: list[list[list[Any]]] = []
    try:
        for worksheet in workbook.worksheets:
            matrix = [
                [cell for cell in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            sheets.append(_rows_from_matrix(matrix))
    finally:
        workbook.close()
    return sheets


def _parse_xls(data: bytes) -> list[list[list[Any]]]:
    try:
        import xlrd
    except ImportError as error:  # pragma: no cover - dependency guarded
        raise ParserError("Legacy .xls support requires the 'xlrd' package.") from error

    try:
        workbook = xlrd.open_workbook(file_contents=data)
    except xlrd.XLRDError as error:
        raise ParserError("The legacy Excel file could not be opened.") from error

    sheets: list[list[list[Any]]] = []
    for sheet in workbook.sheets():
        matrix: list[list[Any]] = []
        for row_index in range(sheet.nrows):
            values = []
            for col_index in range(sheet.ncols):
                cell = sheet.cell(row_index, col_index)
                value: Any = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    year, month, day, *_ = xlrd.xldate_as_tuple(value, workbook.datemode)
                    value = datetime(year, month, day)
                values.append(value)
            matrix.append(values)
        sheets.append(matrix)
    return sheets


class ExcelParser:
    def parse(
        self,
        data: bytes,
        *,
        filename: str = "",
        password: str | None = None,
    ) -> ParsedStatement:
        lowered = (filename or "").lower()

        if lowered.endswith(".xls") and data[:8] != b"PK\x03\x04" and not data.startswith(b"PK"):
            if data.lstrip()[:1] in (b"<", b"{"):
                sheets = [_parse_html_xls(data)]
            else:
                sheets = _parse_xls(data)
        else:
            sheets = _parse_xlsx(data)

        if not sheets:
            raise ParserError("The Excel file contains no sheets.")

        warnings: list[str] = []
        parsed_rows = []
        for index, matrix in enumerate(sheets):
            if not matrix:
                continue
            parsed_rows = rows_to_parsed(matrix, warnings, source_page=index + 1)
            if parsed_rows:
                break

        text_probe = ""
        for matrix in sheets:
            for row in matrix[:20]:
                text_probe += " ".join(str(c) for c in row if c is not None) + "\n"
            if text_probe:
                break

        holder, account = sniff_account_details(text_probe)
        return ParsedStatement(
            rows=parsed_rows,
            bank_name=sniff_bank_name(text_probe),
            account_holder=holder,
            account_number_raw=account,
            warnings=warnings,
            opening_balance=sniff_opening_balance(text_probe),
        )
